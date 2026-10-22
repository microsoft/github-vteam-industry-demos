"""Generate the high-level architecture in SVG, 4K PNG, editable PPTX and Mermaid.

Run from the repository root: python docs\\architecture\\generate.py
No Azure calls, application execution, network requests, or deployment operations.
"""

from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
STEM = "ai-proposal-factory"
WIDTH, HEIGHT = 1920, 1080
NAVY, BLUE, TEAL = "#142D4E", "#0078D4", "#008577"
MUTED, BORDER, PALE, WHITE = "#506680", "#C9D9E8", "#F2F7FC", "#FFFFFF"
FOOTER = ("Built with GitHub Copilot CLI + Squad | "
          "Deployed using Bicep + Azure Developer CLI")
AGENTS = ["Requirements Analyst", "Proposal Writer", "Risk Reviewer"]
ICONS = {
    "app": "02884-icon-service-Worker-Container-App.svg",
    "foundry": "035746832-icon-service-AI-Foundry.svg",
    "search": "10044-icon-service-Cognitive-Search.svg",
    "storage": "10086-icon-service-Storage-Accounts.svg",
}


class Scene:
    def __init__(self):
        self.items = []

    def box(self, x, y, w, h, fill=WHITE, stroke=BORDER, radius=12, name=None):
        self.items.append(dict(kind="box", x=x, y=y, w=w, h=h, fill=fill,
                               stroke=stroke, radius=radius, name=name))

    def text(self, x, y, w, lines, size=28, color=NAVY, bold=False, leading=None,
             center=False):
        if isinstance(lines, str):
            lines = [lines]
        self.items.append(dict(kind="text", x=x, y=y, w=w,
                               h=len(lines) * (leading or size * 1.35),
                               lines=lines, size=size, color=color, bold=bold,
                               leading=leading or size * 1.35, center=center))

    def line(self, points, color=BLUE, width=2.5, name=None):
        self.items.append(dict(kind="line", points=points, color=color,
                               width=width, name=name))

    def icon(self, name, x, y, size=56):
        self.items.append(dict(kind="icon", name=name, x=x, y=y, w=size, h=size))


def build_scene():
    s = Scene()
    s.box(0, 0, WIDTH, HEIGHT, fill="#FAFCFE", stroke=None, radius=0,
          name="background")
    s.text(80, 42, 1200, "AI Proposal Factory", size=44, bold=True)

    s.box(620, 140, 680, 116, stroke=BLUE, name="core-consultant")
    s.text(657, 157, 606, "Professional Services Consultant", size=33, bold=True)
    s.text(657, 205, 606, "Browser", size=27, color=MUTED)
    s.line([(960, 256), (960, 352)], name="flow-https")
    s.text(981, 280, 150, "HTTPS", size=25, color=BLUE)

    s.box(540, 352, 840, 224, stroke=BLUE, name="core-app")
    s.icon("app", 574, 382, 62)
    s.text(659, 381, 685, "Azure Container Apps", size=33, bold=True)
    s.text(659, 434, 685, "AI Proposal Factory", size=35, bold=True)
    s.text(659, 487, 685, "React frontend + Python FastAPI backend", size=27)
    s.text(1110, 529, 235, "PDF export", size=23, color=TEAL)

    # Three equal service cards; all service calls originate at the application.
    for x, key in [(80, "foundry"), (680, "search"), (1280, "storage")]:
        s.line([(x + 280, 644), (x + 280, 744)], name=f"flow-{key}")
    s.line([(720, 576), (720, 644), (360, 644)], name="route-foundry")
    s.line([(960, 576), (960, 644)], name="route-search")
    s.line([(1200, 576), (1200, 644), (1560, 644)], name="route-storage")
    s.text(89, 673, 540, "Agent orchestration", size=25, color=BLUE)
    s.text(704, 673, 510, "Knowledge retrieval", size=25, color=BLUE)
    s.text(1304, 673, 510, "Read / write", size=25, color=BLUE)

    for x, key, title in [
        (80, "foundry", "Microsoft Foundry"),
        (680, "search", "Azure AI Search"),
        (1280, "storage", "Azure Storage"),
    ]:
        s.box(x, 744, 560, 260, name=f"core-{key}")
        s.icon(key, x + 28, 773)
        s.text(x + 106, 783, 426, title, size=33, bold=True)

    for x, lines in [(104, ["Requirements", "Analyst"]),
                     (288, ["Proposal", "Writer"]),
                     (472, ["Risk", "Reviewer"])]:
        s.box(x, 867, 144, 100, fill=PALE, stroke=None, radius=8,
              name="agent-" + "-".join(lines).lower())
        s.text(x + 8, 885, 128, lines, size=20, bold=True, leading=28)
    s.line([(253, 917), (282, 917)], color=TEAL, name="agent-sequence-1")
    s.line([(437, 917), (466, 917)], color=TEAL, name="agent-sequence-2")
    s.text(714, 875, 492,
           ["Organizational knowledge retrieval", "Grounded proposal generation"],
           size=27, leading=43)
    s.text(1314, 885, 492, ["Documents and", "application data"],
           size=30, leading=43)
    s.text(80, 1035, 1760, FOOTER, size=23, color=MUTED, center=True)
    return s


def icon_uri(name):
    data = (ROOT / "icons" / ICONS[name]).read_bytes()
    return "data:image/svg+xml;base64," + base64.b64encode(data).decode("ascii")


def render_svg(scene):
    chunks = [
        f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
        f'width="{WIDTH}" height="{HEIGHT}" viewBox="0 0 {WIDTH} {HEIGHT}" role="img" '
        'aria-labelledby="title desc">',
        '<title id="title">AI Proposal Factory — high-level architecture</title>',
        '<desc id="desc">A consultant uses one application hosted in Azure Container Apps. '
        'The application connects to Microsoft Foundry with three sequential agents, '
        'Azure AI Search and Azure Storage. PDF export is an application output.</desc>',
        "<defs>",
    ]
    for color in [BLUE, TEAL]:
        chunks.append(
            f'<marker id="arrow{color[1:]}" markerWidth="10" markerHeight="10" '
            f'refX="9" refY="5" orient="auto" markerUnits="userSpaceOnUse">'
            f'<path d="M 0 0 L 10 5 L 0 10 Z" fill="{color}"/></marker>'
        )
    chunks.append("</defs>")
    for i, e in enumerate(scene.items):
        kind = e["kind"]
        identifier = f' id="{e["name"]}"' if e.get("name") else ""
        if kind == "box":
            chunks.append(
                f'<rect{identifier} x="{e["x"]}" y="{e["y"]}" width="{e["w"]}" '
                f'height="{e["h"]}" rx="{e["radius"]}" fill="{e["fill"]}" '
                f'stroke="{e["stroke"] or "none"}" stroke-width="1.5"/>')
        elif kind == "text":
            for n, line in enumerate(e["lines"]):
                x = e["x"] + e["w"] / 2 if e["center"] else e["x"]
                anchor = "middle" if e["center"] else "start"
                chunks.append(
                    f'<text id="text-{i}-{n}" x="{x}" y="{e["y"] + n * e["leading"]}" '
                    f'data-max-width="{e["w"]}" dominant-baseline="text-before-edge" '
                    f'text-anchor="{anchor}" '
                    f'font-family="Segoe UI, Arial, sans-serif" font-size="{e["size"]}" '
                    f'font-weight="{600 if e["bold"] else 400}" fill="{e["color"]}">'
                    f'{escape(line)}</text>')
        elif kind == "line":
            points = " ".join(f"{x},{y}" for x, y in e["points"])
            # Routed stems join the final, arrowed service connection without extra arrows.
            arrow = (f'marker-end="url(#arrow{e["color"][1:]})"'
                     if not e["name"].startswith("route-") else "")
            chunks.append(
                f'<polyline{identifier} points="{points}" fill="none" stroke="{e["color"]}" '
                f'stroke-width="{e["width"]}" stroke-linejoin="round" {arrow}/>')
        elif kind == "icon":
            chunks.append(
                f'<image x="{e["x"]}" y="{e["y"]}" width="{e["w"]}" height="{e["h"]}" '
                f'xlink:href="{icon_uri(e["name"])}"/>')
    return "\n".join(chunks + ["</svg>", ""])


def color(value):
    return RGBColor.from_string(value.lstrip("#"))


def unit(value):
    return Inches(value / 144)


def render_pptx(scene, icon_pngs):
    prs = Presentation()
    prs.slide_width, prs.slide_height = unit(WIDTH), unit(HEIGHT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    prs.core_properties.title = "AI Proposal Factory — high-level architecture"
    prs.core_properties.subject = "Application and three supporting Azure services"
    prs.core_properties.author = "AI Proposal Factory"
    for e in scene.items:
        kind = e["kind"]
        if kind == "box":
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE,
                unit(e["x"]), unit(e["y"]), unit(e["w"]), unit(e["h"]))
            if e["radius"]:
                shape.adjustments[0] = min(0.16, e["radius"] / min(e["w"], e["h"]))
            shape.fill.solid()
            shape.fill.fore_color.rgb = color(e["fill"])
            if e["stroke"]:
                shape.line.color.rgb = color(e["stroke"])
                shape.line.width = Pt(0.75)
            else:
                shape.line.fill.background()
        elif kind == "text":
            shape = slide.shapes.add_textbox(
                unit(e["x"]), unit(e["y"]), unit(e["w"]), unit(e["h"] + 3))
            tf = shape.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.TOP
            for n, line in enumerate(e["lines"]):
                p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                p.text = line
                p.alignment = PP_ALIGN.CENTER if e["center"] else PP_ALIGN.LEFT
                p.font.name = "Segoe UI"
                p.font.size = Pt(e["size"] * 0.5)
                p.font.bold = e["bold"]
                p.font.color.rgb = color(e["color"])
                p.line_spacing = Pt(e["leading"] * 0.5)
                p.space_before = p.space_after = Pt(0)
        elif kind == "line":
            points = e["points"]
            for n, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
                shape = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, unit(x1), unit(y1), unit(x2), unit(y2))
                shape.name = e["name"] + f"-{n}"
                shape.line.color.rgb = color(e["color"])
                shape.line.width = Pt(e["width"] * 0.5)
                if n == len(points) - 2 and not e["name"].startswith("route-"):
                    arrow = OxmlElement("a:tailEnd")
                    arrow.set("type", "triangle")
                    shape._element.spPr.get_or_add_ln().append(arrow)
            continue
        elif kind == "icon":
            shape = slide.shapes.add_picture(
                BytesIO(icon_pngs[e["name"]]), unit(e["x"]), unit(e["y"]),
                unit(e["w"]), unit(e["h"]))
        if e.get("name"):
            shape.name = e["name"]
    slide.notes_slide.notes_text_frame.text = (
        "Source: docs/architecture/generate.py. Cards, text and connectors are native "
        "editable objects; only the four official Azure icons are pictures. "
        "See docs/architecture/README.md for retained technical evidence and icon terms. "
        "Foundry arrows indicate application-orchestrated execution order."
    )
    prs.save(ROOT / f"{STEM}.pptx")


def render_mermaid():
    return f'''flowchart TB
    %%{{init: {{"theme":"base","htmlLabels":false,"fontFamily":"Segoe UI","themeVariables":{{"fontFamily":"Segoe UI","primaryTextColor":"#142d4e","lineColor":"#0078d4","edgeLabelBackground":"#fafcfe"}},"flowchart":{{"htmlLabels":false,"wrappingWidth":1600,"curve":"linear"}}}}}}%%
    consultant["Professional Services Consultant<br/>Browser"]
    app["Azure Container Apps<br/>AI Proposal Factory<br/>React frontend + Python FastAPI backend<br/>PDF export"]
    subgraph foundry["Microsoft Foundry"]
        direction LR
        analyst["Requirements<br/>Analyst"]
        writer["Proposal<br/>Writer"]
        reviewer["Risk<br/>Reviewer"]
        analyst --> writer --> reviewer
    end
    search["Azure AI Search<br/>Organizational knowledge retrieval<br/>Grounded proposal generation"]
    storage["Azure Storage<br/>Documents and application data"]
    consultant -->|HTTPS| app
    app -->|Agent orchestration| foundry
    app -->|Knowledge retrieval| search
    app -->|Read / write| storage
    footer["{FOOTER}"]
    storage ~~~ footer

    classDef core fill:#ffffff,stroke:#0078d4,color:#142d4e
    classDef agent fill:#f2f7fc,stroke:#c9d9e8,color:#142d4e
    classDef annotation fill:transparent,stroke:transparent,color:#506680,font-size:14px
    class consultant,app,search,storage core
    class analyst,writer,reviewer agent
    class footer annotation
    style foundry fill:#ffffff,stroke:#c9d9e8,color:#142d4e
'''


def main():
    scene = build_scene()
    svg = ROOT / f"{STEM}.svg"
    svg.write_text(render_svg(scene), encoding="utf-8")
    (ROOT / f"{STEM}.mmd").write_text(render_mermaid(), encoding="utf-8")
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": WIDTH, "height": HEIGHT},
                                device_scale_factor=2)
        page.route("**/*", lambda route: route.continue_()
                   if route.request.url.startswith("file:") else route.abort())
        page.goto(svg.as_uri(), wait_until="load")
        page.evaluate("document.fonts.ready")
        overflows = page.evaluate("""() => [...document.querySelectorAll('text[data-max-width]')]
            .filter(e => {
                const b=e.getBBox();
                return b.width > Number(e.dataset.maxWidth) + 1 ||
                    b.x < 0 || b.y < 0 || b.x+b.width > 1920 || b.y+b.height > 1080;
            }).map(e => e.textContent)""")
        if overflows:
            raise RuntimeError(f"Text exceeds allocated or canvas bounds: {overflows}")
        page.screenshot(path=str(ROOT / f"{STEM}.png"))
        page.close()
        page = browser.new_page(viewport={"width": 96, "height": 96}, device_scale_factor=2)
        icon_pngs = {}
        for name in ICONS:
            page.set_content(
                '<body style="margin:0;background:transparent">'
                f'<img width="96" height="96" src="{icon_uri(name)}"></body>')
            page.wait_for_function("document.querySelector('img').complete")
            icon_pngs[name] = page.locator("img").screenshot(omit_background=True)
        browser.close()
    render_pptx(scene, icon_pngs)
    with Image.open(ROOT / f"{STEM}.png") as image:
        assert image.size == (3840, 2160)
    print("Generated SVG, 3840x2160 PNG, one native editable PPTX slide and Mermaid.")
    print("Browser-measured text bounds passed. No Azure or network calls performed.")


if __name__ == "__main__":
    main()
