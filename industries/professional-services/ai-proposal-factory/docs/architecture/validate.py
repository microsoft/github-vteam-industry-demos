"""Focused, offline structure, cross-format, editability and text-bounds checks."""

import base64
import re
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from generate import AGENTS, FOOTER, build_scene, render_mermaid, render_svg, unit

ROOT = Path(__file__).resolve().parent
STEM = "ai-proposal-factory"
SVG = "{http://www.w3.org/2000/svg}"
XLINK = "{http://www.w3.org/1999/xlink}"
PPT = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
DRAW = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

scene = build_scene()
core = {e["name"]: e for e in scene.items if
        e["kind"] == "box" and (e.get("name") or "").startswith("core-")}
assert set(core) == {"core-consultant", "core-app", "core-foundry",
                     "core-search", "core-storage"}
services = [core["core-" + key] for key in ["foundry", "search", "storage"]]
assert len({(e["w"], e["h"], e["y"]) for e in services}) == 1
assert core["core-consultant"]["y"] < core["core-app"]["y"] < services[0]["y"]
agents = [e for e in scene.items if
          e["kind"] == "box" and (e.get("name") or "").startswith("agent-")]
assert len(agents) == 3
foundry = core["core-foundry"]
for agent in agents:
    assert foundry["x"] < agent["x"] < agent["x"] + agent["w"] < foundry["x"] + foundry["w"]
    assert foundry["y"] < agent["y"] < agent["y"] + agent["h"] < foundry["y"] + foundry["h"]
assert [e["name"] for e in scene.items if e["kind"] == "line"] == [
    "flow-https", "flow-foundry", "flow-search", "flow-storage",
    "route-foundry", "route-search", "route-storage",
    "agent-sequence-1", "agent-sequence-2"]
print("PASS scene: exactly five core elements, three sequential Foundry agents, "
      "equal service cards and four application/client connections.")

with Image.open(ROOT / f"{STEM}.png") as image:
    assert image.size == (3840, 2160), image.size
    image.verify()
print("PASS PNG: valid 3840x2160 (16:9).")

assert (ROOT / f"{STEM}.svg").read_text("utf-8") == render_svg(scene)
assert (ROOT / f"{STEM}.mmd").read_text("utf-8") == render_mermaid()
svg = ET.parse(ROOT / f"{STEM}.svg").getroot()
assert svg.attrib["viewBox"] == "0 0 1920 1080"
texts = [e.text or "" for e in svg.iter(SVG + "text")]
assert len(list(svg.iter(SVG + "image"))) == 4
for e in svg.iter(SVG + "image"):
    href = e.attrib[XLINK + "href"]
    assert href.startswith("data:image/svg+xml;base64,")
    assert ET.fromstring(base64.b64decode(href.split(",", 1)[1])).tag == SVG + "svg"
assert not list(svg.iter(SVG + "script")) and not list(svg.iter(SVG + "foreignObject"))
print(f"PASS SVG: self-contained, {len(texts)} editable text lines; four official icons.")

prs = Presentation(ROOT / f"{STEM}.pptx")
assert len(prs.slides) == 1
assert abs(prs.slide_width / prs.slide_height - 16 / 9) < 0.00001
shapes = prs.slides[0].shapes
for shape in shapes:
    assert shape.left >= 0 and shape.top >= 0, shape.name
    assert shape.left + shape.width <= prs.slide_width + 1, shape.name
    assert shape.top + shape.height <= prs.slide_height + 1, shape.name
counts = Counter(shape.shape_type for shape in shapes)
expected_text = [e for e in scene.items if e["kind"] == "text"]
expected_boxes = [e for e in scene.items if e["kind"] == "box"]
expected_segments = sum(len(e["points"]) - 1 for e in scene.items if e["kind"] == "line")
assert counts[MSO_SHAPE_TYPE.TEXT_BOX] == len(expected_text), counts
assert counts[MSO_SHAPE_TYPE.AUTO_SHAPE] == len(expected_boxes), counts
assert counts[MSO_SHAPE_TYPE.LINE] == expected_segments, counts
assert counts[MSO_SHAPE_TYPE.PICTURE] == 4, counts
native_text = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
for shape, item in zip(native_text, expected_text):
    assert shape.text == "\n".join(item["lines"])
    assert (shape.left, shape.top, shape.width) == (
        unit(item["x"]), unit(item["y"]), unit(item["w"]))
    for p in shape.text_frame.paragraphs:
        assert abs(p.font.size.pt - item["size"] * 0.5) < 0.01
assert {s.name for s in shapes if s.name.startswith("core-")} == set(core)
with ZipFile(ROOT / f"{STEM}.pptx") as archive:
    assert archive.testzip() is None
    slide = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    assert len(list(slide.iter(PPT + "cxnSp"))) == expected_segments
    assert len(list(slide.iter(DRAW + "tailEnd"))) == 6
    assert not any(name.startswith("ppt/embeddings/") for name in archive.namelist())
print("PASS PPTX: one editable slide; "
      f"{counts[MSO_SHAPE_TYPE.TEXT_BOX]} text boxes, "
      f"{counts[MSO_SHAPE_TYPE.AUTO_SHAPE]} shapes, {expected_segments} native connectors, "
      "four icon-only pictures; bounds and font sizes match SVG.")

mermaid = (ROOT / f"{STEM}.mmd").read_text("utf-8")
assert mermaid.startswith("flowchart TB\n")
assert mermaid.count("subgraph ") == 1 and 'subgraph foundry["Microsoft Foundry"]' in mermaid
assert 'class consultant,app,search,storage core' in mermaid
assert 'class analyst,writer,reviewer agent' in mermaid
assert 'class footer annotation' in mermaid
assert 'analyst --> writer --> reviewer' in mermaid
assert re.findall(r"^\s*(\w+) -->\|([^|]+)\| (\w+)$", mermaid, re.M) == [
    ("consultant", "HTTPS", "app"), ("app", "Agent orchestration", "foundry"),
    ("app", "Knowledge retrieval", "search"), ("app", "Read / write", "storage")]
required = [
    "Professional Services Consultant", "Browser", "Azure Container Apps",
    "AI Proposal Factory", "React frontend + Python FastAPI backend",
    "Microsoft Foundry", *AGENTS, "Azure AI Search",
    "Organizational knowledge retrieval", "Grounded proposal generation",
    "Azure Storage", "Documents and application data", "PDF export",
]
contents = {
    "SVG": " ".join(texts),
    "PPTX": " ".join(s.text.replace("\n", " ") for s in native_text),
    "Mermaid": mermaid.replace("<br/>", " "),
}
for fmt, content in contents.items():
    for label in required:
        assert label in content, f"{fmt} missing {label}"
    assert content.count(FOOTER) == 1, fmt
    assert content.count("PDF export") == 1, fmt
    assert not re.search(
        r"rg-|ca-|cae-|proj-|stprof|gpt-|GlobalStandard|requirements-analyst:|"
        r"proposal-writer:|risk-reviewer:|managed identity|RBAC|Log Analytics|"
        r"Environment|port 8000|/api/|Blob|Table|Sweden|swedencentral|"
        r"202[0-9]-|verification|TypeScript|Markdown|Lead / Architect|"
        r"Full-Stack|Azure \+ AI Engineer|ACR|DefaultAzureCredential",
        content, re.I), f"{fmt} contains excluded detail"
print("PASS cross-format labels, app-only PDF annotation, exact footer and detail exclusions.")

# Measure actual glyph bounds in Chromium, not an estimated character count.
with sync_playwright() as p:
    browser = p.chromium.launch()
    page = browser.new_page()
    page.route("**/*", lambda route: route.continue_()
               if route.request.url.startswith("file:") else route.abort())
    page.goto((ROOT / f"{STEM}.svg").as_uri())
    page.evaluate("document.fonts.ready")
    issues = page.evaluate("""() => {
        const texts=[...document.querySelectorAll('text')];
        const cards=[...document.querySelectorAll('rect[id^="core-"],rect[id^="agent-"]')];
        const problems=[];
        const intersects=(a,b) => a.x < b.x+b.width && a.x+a.width > b.x &&
            a.y < b.y+b.height && a.y+a.height > b.y;
        texts.forEach((t,i) => {
            const a=t.getBBox();
            if(a.width>Number(t.dataset.maxWidth)+1 || a.x<0 || a.y<0 ||
               a.x+a.width>1920 || a.y+a.height>1080) problems.push(t.textContent);
            texts.slice(i+1).forEach(u => {
                if(intersects(a,u.getBBox())) problems.push('Text overlap: '+t.textContent);
            });
            cards.forEach(c => {
                const b=c.getBBox();
                if(intersects(a,b) && !(a.x>=b.x && a.y>=b.y &&
                   a.x+a.width<=b.x+b.width && a.y+a.height<=b.y+b.height))
                    problems.push('Text crosses card: '+t.textContent);
            });
        });
        return problems;
    }""")
    browser.close()
assert not issues, issues
print("PASS browser text bounds: no clipping, text overlap or card-border crossings.")

readme = (ROOT.parent.parent / "README.md").read_text("utf-8")
for ext in ["png", "svg", "pptx", "mmd"]:
    assert f"docs/architecture/{STEM}.{ext}" in readme
print("PASS README links. Run validate_mermaid.py for local Mermaid parse/render checks.")
